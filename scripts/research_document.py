"""Designed source-backed documents. Input is data only; no network or code evaluation."""
from pathlib import Path
from html import escape
import json

PALETTES={'editorial':('#b74325','#fff8f2'),'navy':('#183850','#f0f5fa'),'forest':('#235d46','#f2f7f1')}

def export_document(folder):
    from docx import Document
    from docx.shared import Inches,Pt,RGBColor
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,KeepTogether
    from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
    from reportlab.lib.colors import HexColor
    from reportlab.lib.enums import TA_LEFT
    data=json.loads((folder/'document.json').read_text(encoding='utf-8'))
    accent,paper=PALETTES.get(data['theme'],PALETTES['editorial'])
    def citation(item):return ' '+''.join('['+s+']' for s in item['sourceIds']) if item['sourceIds'] else ''
    md=['# '+data['title'],'',data['subtitle'].title()+' · '+data['audience'],'','## Purpose','',data['purpose'],'']
    doc=Document();section=doc.sections[0];section.top_margin=Inches(.85);section.bottom_margin=Inches(.8);section.left_margin=section.right_margin=Inches(.9)
    normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10.5);normal.paragraph_format.space_after=Pt(9)
    for style,size in [('Title',32),('Heading 1',17),('Heading 2',13)]:doc.styles[style].font.name='Calibri';doc.styles[style].font.size=Pt(size);doc.styles[style].font.color.rgb=RGBColor.from_string(accent[1:])
    doc.core_properties.title=data['title'];doc.core_properties.author='Command Deck';doc.add_heading(data['title'],0);doc.add_paragraph(data['subtitle'].title()+' · Prepared for '+data['audience']);doc.add_heading('Purpose',1);doc.add_paragraph(data['purpose'])
    header=section.header.paragraphs[0];header.text='COMMAND DECK  /  RESEARCH';header.style='Caption'
    section.footer.paragraphs[0].text='Source-backed material · '+data['createdAt'][:10]
    styles=getSampleStyleSheet();styles.add(ParagraphStyle(name='DocTitle',fontName='Times-Roman',fontSize=32,leading=36,textColor=HexColor(accent),spaceAfter=18));styles.add(ParagraphStyle(name='DocText',fontName='Helvetica',fontSize=10,leading=15,spaceAfter=11));styles.add(ParagraphStyle(name='DocHeading',fontName='Helvetica-Bold',fontSize=15,leading=20,textColor=HexColor(accent),spaceBefore=20,spaceAfter=9));styles.add(ParagraphStyle(name='DocSmall',fontName='Helvetica',fontSize=8,leading=12,textColor=HexColor('#5d6470'),spaceAfter=8))
    story=[Paragraph(escape(data['title']),styles['DocTitle']),Paragraph(escape(data['subtitle'].title())+' · Prepared for '+escape(data['audience']),styles['DocText']),Spacer(1,12),Paragraph('Purpose',styles['DocHeading']),Paragraph(escape(data['purpose']),styles['DocText'])]
    html_sections=[]
    for block in data['sections']:
        md+=['## '+block['heading'],''];doc.add_heading(block['heading'],1);story.append(Paragraph(escape(block['heading']),styles['DocHeading']));paragraphs=[]
        for item in block['paragraphs']:
            text=item['text']+citation(item);md+=[text,''];doc.add_paragraph(text);story.append(Paragraph(escape(text),styles['DocText']));paragraphs.append('<p>'+escape(item['text'])+'<sup>'+escape(citation(item))+'</sup></p>')
        html_sections.append('<section><h2>'+escape(block['heading'])+'</h2>'+''.join(paragraphs)+'</section>')
    md+=['## Sources',''];doc.add_heading('Sources',1);story.append(Paragraph('Sources',styles['DocHeading']));source_html=[]
    for s in data['sources']:
        text='['+s['id']+'] '+s['title']+' — '+s['url']+' · retrieved '+s['at'][:10];md+=[text,''];doc.add_paragraph(text,'Caption');story.append(Paragraph(escape(text),styles['DocSmall']));source_html.append('<li><strong>['+escape(s['id'])+'] '+escape(s['title'])+'</strong><br><a href="'+escape(s['url'],quote=True)+'">'+escape(s['url'])+'</a><br>Retrieved '+escape(s['at'][:10])+'</li>')
    md+=['## Collection notes',''];doc.add_heading('Collection notes',1);story.append(Paragraph('Collection notes',styles['DocHeading']))
    for note in data['notes']:md+=['- '+note];doc.add_paragraph(note);story.append(Paragraph(escape(note),styles['DocSmall']))
    doc.save(folder/'document.docx');(folder/'document.md').write_text('\n'.join(md),encoding='utf-8')
    def page(canvas,document):
        canvas.setStrokeColor(HexColor(accent));canvas.setLineWidth(2);canvas.line(54,810,541,810);canvas.setFont('Helvetica',8);canvas.setFillColor(HexColor('#5d6470'));canvas.drawString(54,30,'COMMAND DECK / SOURCE-BACKED RESEARCH');canvas.drawRightString(541,30,str(document.page))
    SimpleDocTemplate(str(folder/'document.pdf'),pagesize=(595.28,841.89),title=data['title'],author='Command Deck',topMargin=58,bottomMargin=52,leftMargin=54,rightMargin=54).build(story,onFirstPage=page,onLaterPages=page)
    html='''<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width"><title>'''+escape(data['title'])+'''</title><style>body{margin:0;background:'''+paper+''';color:#202832;font:16px/1.75 Georgia,serif}main{max-width:780px;margin:48px auto;background:white;padding:64px}header{border-bottom:2px solid '''+accent+''';padding-bottom:32px}h1{font-size:46px;line-height:1.12;font-weight:400;color:'''+accent+''';margin:16px 0 24px}h2{font:600 21px/1.4 Arial,sans-serif;margin:36px 0 14px;color:'''+accent+'''}small,li{font:13px/1.6 Arial,sans-serif}sup{font:11px Arial;color:'''+accent+'''}li{margin:18px 0;overflow-wrap:anywhere}a{color:inherit}@media(max-width:600px){main{margin:0;padding:28px}h1{font-size:34px}}@media print{body{background:white}main{margin:0;padding:0}section{break-inside:avoid}}</style></head><body><main><header><h1>'''+escape(data['title'])+'</h1><p>'+escape(data['purpose'])+'</p><small>'+escape(data['subtitle'].title())+' · Prepared for '+escape(data['audience'])+'</small></header>'+''.join(html_sections)+'<section><h2>Sources</h2><ol>'+''.join(source_html)+'</ol></section><section><h2>Collection notes</h2>'+''.join('<p>'+escape(n)+'</p>' for n in data['notes'])+'</section></main></body></html>'
    (folder/'document.html').write_text(html,encoding='utf-8')
    return ['document.docx','document.pdf','document.html','document.md','document.json']
