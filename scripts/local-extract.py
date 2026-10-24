"""Bounded local extraction using the already installed document runtime. No network calls."""
import argparse,json,sys,subprocess,tempfile,zipfile
from pathlib import Path
def ocr(file):
    script=Path(__file__).with_name('local-ocr.ps1')
    result=subprocess.run(['powershell.exe','-NoProfile','-ExecutionPolicy','Bypass','-File',str(script),'-ImagePath',str(file)],capture_output=True,text=True,encoding='utf-8',errors='replace',timeout=60)
    if result.returncode: raise RuntimeError('Windows OCR unavailable. Install an OCR language or use Docling after freeing disk space.')
    return result.stdout.strip()
def extract(file):
    text=[];warnings=[];suffix=file.suffix.lower()
    if suffix in ['.docx','.pptx']:
        with zipfile.ZipFile(file) as archive:
            if sum(x.file_size for x in archive.infolist())>100000000:raise RuntimeError('Expanded document exceeds the 100 MB limit.')
    if suffix=='.docx':
        from docx import Document
        doc=Document(file);text=[p.text for p in doc.paragraphs]+[' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows]
    elif suffix=='.pptx':
        from pptx import Presentation
        doc=Presentation(file);text=[shape.text for slide in list(doc.slides)[:100] for shape in slide.shapes if shape.has_text_frame]
        if len(doc.slides)>100:warnings.append('Only the first 100 slides were extracted.')
    elif suffix=='.pdf':
        from pypdf import PdfReader
        doc=PdfReader(file);limit=min(len(doc.pages),100)
        for index in range(limit):
            page=doc.pages[index];value=page.extract_text() or ''
            if not value.strip():
                import pypdfium2
                pdf=pypdfium2.PdfDocument(file)
                try:
                    with tempfile.TemporaryDirectory(prefix='deck-ocr-') as temp:
                        image=Path(temp)/'page.png';pdf[index].render(scale=min(1.25,2000/max(pdf[index].get_size()))).to_pil().save(image);value=ocr(image)
                finally:pdf.close()
            text.append(f'## Page {index+1}\n\n{value}')
        if len(doc.pages)>100:warnings.append('Only the first 100 pages were extracted.')
    else:text=[ocr(file)]
    value='\n\n'.join(text).strip()
    if not value:raise RuntimeError('No text was extracted; review the original document.')
    if len(value)>200000:value=value[:200000];warnings.append('Text truncated at 200,000 characters; original retained.')
    return {'markdown':value,'warnings':warnings,'engine':'Native local extraction (Windows OCR for scans)','layout':'Text and tables; original visual layout is not reproduced.'}
if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('file',nargs='?');parser.add_argument('--check',action='store_true');args=parser.parse_args()
    try:
        if args.check:
            import docx,pypdf,pypdfium2,pptx
            print(json.dumps({'ready':True}));sys.exit(0)
        print(json.dumps(extract(Path(args.file).resolve()),ensure_ascii=True))
    except Exception as error:
        print(json.dumps({'error':str(error)}));sys.exit(1)

